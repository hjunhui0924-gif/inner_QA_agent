"""SQLite memory and knowledge base helpers."""

from __future__ import annotations

import asyncio
import csv
import json
import os
import tempfile
import threading
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Protocol

import aiosqlite
from filelock import FileLock
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

try:
    from langchain_chroma import Chroma
except ImportError:  # pragma: no cover
    try:
        from langchain_community.vectorstores import Chroma
    except ImportError:  # pragma: no cover
        Chroma = None  # type: ignore[assignment]

from backend.config import settings
from backend.knowledge.chunking import split_text
from backend.knowledge.deduplication import content_fingerprint, find_duplicate
from backend.knowledge.embeddings import EmbeddingProfile, create_embeddings
from backend.retrieval.engine import RetrievalConfig, RetrievalEngine, tokenize
from backend.retrieval.reranker import DashScopeReranker


class VectorStoreLike(Protocol):
    """Minimal vector store interface used by the app."""

    def add_documents(
        self,
        documents: list[Document],
        **kwargs: Any,
    ) -> Any:
        """Add documents to the index."""

    def similarity_search(self, query: str, k: int = 4) -> list[Document]:
        """Run a similarity search."""

    def get(self, **kwargs: Any) -> dict[str, Any]:
        """Return persisted vector-store records."""

    def delete(self, ids: list[str]) -> Any:
        """Delete persisted records by stable ID."""


_VECTORSTORE: VectorStoreLike | None = None
_RETRIEVER: RetrievalEngine | None = None
_RETRIEVER_KB_VERSION: tuple[int, int] | None = None
_KNOWLEDGE_WRITE_LOCK = threading.RLock()


@dataclass(frozen=True)
class ExtractedDocument:
    """Parsed text plus source locations that must survive indexing."""

    content: str
    segments: list[dict[str, Any]]


class _LocalVectorStore:
    """In-memory fallback when Chroma dependencies are unavailable."""

    def __init__(self, embeddings: Embeddings, documents: list[Document]) -> None:
        self._embeddings = embeddings
        self._documents = list(documents)
        self._ids = [
            str(document.metadata.get("chunk_id", index))
            for index, document in enumerate(documents)
        ]
        self._vectors = embeddings.embed_documents(
            [document.page_content for document in documents]
        )

    def add_documents(
        self,
        documents: list[Document],
        **kwargs: Any,
    ) -> None:
        ids = [str(item) for item in kwargs.get("ids", [])]
        if len(ids) != len(documents):
            ids = [
                str(document.metadata.get("chunk_id", len(self._ids) + index))
                for index, document in enumerate(documents)
            ]
        vectors = self._embeddings.embed_documents(
            [document.page_content for document in documents]
        )
        for item_id, document, vector in zip(ids, documents, vectors, strict=True):
            if item_id in self._ids:
                index = self._ids.index(item_id)
                self._documents[index] = document
                self._vectors[index] = vector
            else:
                self._ids.append(item_id)
                self._documents.append(document)
                self._vectors.append(vector)

    def similarity_search(self, query: str, k: int = 4) -> list[Document]:
        query_vector = self._embeddings.embed_query(query)
        scored: list[tuple[float, Document]] = []
        for vector, document in zip(self._vectors, self._documents, strict=False):
            score = sum(a * b for a, b in zip(query_vector, vector, strict=False))
            scored.append((score, document))
        scored.sort(key=lambda item: item[0], reverse=True)
        return [document for _, document in scored[:k]]

    def get(self, **kwargs: Any) -> dict[str, Any]:
        return {
            "ids": list(self._ids),
            "documents": [document.page_content for document in self._documents],
            "metadatas": [dict(document.metadata) for document in self._documents],
        }

    def delete(self, ids: list[str]) -> None:
        rejected = set(ids)
        retained = [
            (item_id, document, vector)
            for item_id, document, vector in zip(
                self._ids,
                self._documents,
                self._vectors,
                strict=True,
            )
            if item_id not in rejected
        ]
        self._ids = [item[0] for item in retained]
        self._documents = [item[1] for item in retained]
        self._vectors = [item[2] for item in retained]


def ensure_data_directories() -> None:
    """Ensure local data directories exist."""

    Path(settings.chroma_persist_dir).mkdir(parents=True, exist_ok=True)
    Path(settings.upload_dir).mkdir(parents=True, exist_ok=True)
    Path(settings.sqlite_db_path).parent.mkdir(parents=True, exist_ok=True)
    Path(settings.knowledge_base_path).parent.mkdir(parents=True, exist_ok=True)


def set_vectorstore(vectorstore: VectorStoreLike) -> None:
    """Register the shared vector store instance."""

    global _VECTORSTORE
    _VECTORSTORE = vectorstore


def set_retriever(retriever: RetrievalEngine) -> None:
    """Register the shared hybrid retrieval engine."""

    global _RETRIEVER, _RETRIEVER_KB_VERSION
    _RETRIEVER = retriever
    _RETRIEVER_KB_VERSION = _knowledge_base_version()


def get_vectorstore() -> VectorStoreLike:
    """Return the shared vector store instance."""

    if _VECTORSTORE is None:
        raise RuntimeError("Vector store has not been initialized yet.")
    return _VECTORSTORE


def get_retriever() -> RetrievalEngine:
    """Return the initialized hybrid retrieval engine."""

    _refresh_retriever_if_stale()
    if _RETRIEVER is None:
        raise RuntimeError("Retrieval engine has not been initialized yet.")
    return _RETRIEVER


def _knowledge_base_version() -> tuple[int, int] | None:
    path = Path(settings.knowledge_base_path)
    if not path.exists():
        return None
    stat = path.stat()
    return (stat.st_mtime_ns, stat.st_size)


def _refresh_retriever_if_stale() -> None:
    global _RETRIEVER, _RETRIEVER_KB_VERSION
    current_version = _knowledge_base_version()
    if _RETRIEVER is None or current_version == _RETRIEVER_KB_VERSION:
        return
    with _KNOWLEDGE_WRITE_LOCK:
        current_version = _knowledge_base_version()
        if current_version == _RETRIEVER_KB_VERSION:
            return
        _RETRIEVER = _build_retrieval_engine(get_vectorstore())
        _RETRIEVER_KB_VERSION = current_version


async def ensure_user_memory_db(db_path: str | Path) -> None:
    """Create the application tables when needed."""

    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    async with aiosqlite.connect(path) as db:
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS chat_sessions (
                user_id TEXT NOT NULL,
                session_id TEXT NOT NULL,
                title TEXT NOT NULL DEFAULT '',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (user_id, session_id)
            )
            """
        )
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS chat_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id TEXT NOT NULL,
                session_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                citations_json TEXT NOT NULL DEFAULT '[]',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        cursor = await db.execute("PRAGMA table_info(chat_messages)")
        columns = {str(row[1]) for row in await cursor.fetchall()}
        await cursor.close()
        if "citations_json" not in columns:
            await db.execute(
                "ALTER TABLE chat_messages "
                "ADD COLUMN citations_json TEXT NOT NULL DEFAULT '[]'"
            )
        await db.commit()


def _derive_session_title(message: str) -> str:
    """Create a short session title from the first user message."""

    normalized = " ".join(message.strip().split())
    if not normalized:
        return "新会话"
    return normalized[:24]


async def upsert_chat_session(
    user_id: str,
    session_id: str,
    title: str = "",
) -> None:
    """Create or update one chat session record."""

    clean_title = title.strip()
    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        await db.execute(
            """
            INSERT INTO chat_sessions (user_id, session_id, title, updated_at)
            VALUES (?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, session_id) DO UPDATE SET
                title = CASE
                    WHEN chat_sessions.title IN ('', '新会话') AND excluded.title <> ''
                    THEN excluded.title
                    ELSE chat_sessions.title
                END,
                updated_at = CURRENT_TIMESTAMP
            """,
            (user_id, session_id, clean_title),
        )
        await db.commit()


async def append_chat_message(
    user_id: str,
    session_id: str,
    role: str,
    content: str,
    citations: list[dict[str, Any]] | None = None,
) -> None:
    """Append one chat message and touch its session."""

    clean_content = content.strip()
    if not clean_content:
        return

    await upsert_chat_session(
        user_id=user_id,
        session_id=session_id,
        title=_derive_session_title(clean_content) if role == "user" else "",
    )
    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        await db.execute(
            """
            INSERT INTO chat_messages (
                user_id, session_id, role, content, citations_json
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                user_id,
                session_id,
                role,
                clean_content,
                json.dumps(citations or [], ensure_ascii=False),
            ),
        )
        await db.execute(
            """
            UPDATE chat_sessions
            SET updated_at = CURRENT_TIMESTAMP
            WHERE user_id = ? AND session_id = ?
            """,
            (user_id, session_id),
        )
        await db.commit()


async def list_chat_sessions(user_id: str) -> list[dict[str, Any]]:
    """Return chat sessions for one user ordered by recent activity."""

    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        cursor = await db.execute(
            """
            SELECT s.session_id, s.title, s.created_at, s.updated_at,
                   COALESCE(
                       (
                           SELECT m.content
                           FROM chat_messages m
                           WHERE m.user_id = s.user_id
                             AND m.session_id = s.session_id
                           ORDER BY m.id DESC
                           LIMIT 1
                       ),
                       ''
                   ) AS last_message
            FROM chat_sessions s
            WHERE s.user_id = ?
            ORDER BY s.updated_at DESC
            """,
            (user_id,),
        )
        rows = await cursor.fetchall()
        await cursor.close()

    return [
        {
            "session_id": row[0],
            "title": row[1] or "新会话",
            "created_at": row[2],
            "updated_at": row[3],
            "last_message": row[4],
        }
        for row in rows
    ]


async def list_session_title_candidates(user_id: str) -> list[dict[str, str]]:
    """Return sessions whose title still mirrors their first user question."""

    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        cursor = await db.execute(
            """
            SELECT s.session_id, s.title,
                   COALESCE((
                       SELECT m.content FROM chat_messages m
                       WHERE m.user_id = s.user_id
                         AND m.session_id = s.session_id
                         AND m.role = 'user'
                       ORDER BY m.id ASC LIMIT 1
                   ), '')
            FROM chat_sessions s
            WHERE s.user_id = ?
              AND EXISTS (
                  SELECT 1 FROM chat_messages raw_title
                  WHERE raw_title.user_id = s.user_id
                    AND raw_title.session_id = s.session_id
                    AND raw_title.role = 'user'
                    AND substr(trim(raw_title.content), 1, 24) = trim(s.title)
              )
            """,
            (user_id,),
        )
        rows = await cursor.fetchall()
        await cursor.close()
    candidates: list[dict[str, str]] = []
    for session_id, title, first_question in rows:
        clean_question = str(first_question).strip()
        clean_title = str(title).strip()
        if clean_question and clean_title:
            candidates.append(
                {"session_id": str(session_id), "first_question": clean_question}
            )
    return candidates


async def update_chat_session_title(
    user_id: str,
    session_id: str,
    title: str,
) -> None:
    """Replace one session title with a generated summary."""

    clean_title = " ".join(title.strip().split())[:24]
    if not clean_title:
        return
    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        await db.execute(
            "UPDATE chat_sessions SET title = ? WHERE user_id = ? AND session_id = ?",
            (clean_title, user_id, session_id),
        )
        await db.commit()


async def load_chat_messages(user_id: str, session_id: str) -> list[dict[str, Any]]:
    """Return chat messages for one session."""

    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        cursor = await db.execute(
            """
            SELECT role, content, citations_json
            FROM chat_messages
            WHERE user_id = ? AND session_id = ?
            ORDER BY id ASC
            """,
            (user_id, session_id),
        )
        rows = await cursor.fetchall()
        await cursor.close()

    result: list[dict[str, Any]] = []
    for role, content, citations_json in rows:
        try:
            citations = json.loads(citations_json or "[]")
        except (json.JSONDecodeError, TypeError):
            citations = []
        result.append(
            {
                "role": role,
                "content": content,
                "citations": citations if isinstance(citations, list) else [],
            }
        )
    return result


async def delete_chat_session(user_id: str, session_id: str) -> None:
    """Delete one session and all of its messages."""

    async with aiosqlite.connect(settings.sqlite_db_path) as db:
        await db.execute(
            """
            DELETE FROM chat_messages
            WHERE user_id = ? AND session_id = ?
            """,
            (user_id, session_id),
        )
        await db.execute(
            """
            DELETE FROM chat_sessions
            WHERE user_id = ? AND session_id = ?
            """,
            (user_id, session_id),
        )
        await db.commit()


def _load_json_file(path: Path) -> list[dict[str, Any]]:
    """Load JSON array content from disk."""

    if not path.exists():
        return []
    with path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    return data if isinstance(data, list) else []


def _save_json_file(path: Path, data: list[dict[str, Any]]) -> None:
    """Atomically persist JSON so interruption cannot truncate the knowledge base."""

    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".tmp",
        dir=path.parent,
        text=True,
    )
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary_name, path)
    except Exception:
        try:
            Path(temporary_name).unlink(missing_ok=True)
        finally:
            raise


def _normalize_segments(value: Any) -> list[dict[str, Any]]:
    """Validate persisted parser output before it reaches chunk metadata."""

    if not isinstance(value, list):
        return []
    normalized: list[dict[str, Any]] = []
    for item in value:
        if not isinstance(item, dict):
            continue
        text = str(item.get("text", "")).strip()
        if not text:
            continue
        segment: dict[str, Any] = {"text": text}
        page = item.get("page")
        try:
            page_number = int(page) if page is not None and page != "" else None
        except (TypeError, ValueError):
            page_number = None
        if page_number is not None and page_number > 0:
            segment["page"] = page_number
        section = str(item.get("section", "")).strip()
        if section:
            segment["section"] = section[:300]
        normalized.append(segment)
    return normalized


def _load_knowledge_base_records() -> list[dict[str, Any]]:
    """Load the raw knowledge base JSON records."""

    path = Path(settings.knowledge_base_path)
    records = _load_json_file(path)
    normalized: list[dict[str, Any]] = []
    for record in records:
        if not isinstance(record, dict):
            continue
        title = str(record.get("title", "")).strip()
        content = str(record.get("content", "")).strip()
        source = str(record.get("source", "seed")).strip() or "seed"
        if title and content:
            normalized.append(
                {
                    "title": title,
                    "content": content,
                    "source": source,
                    "original_filename": str(
                        record.get("original_filename", "")
                    ).strip(),
                    "segments": _normalize_segments(record.get("segments", [])),
                    # Recompute so manually migrated/edited legacy JSON cannot
                    # point at vectors created from different content.
                    "content_fingerprint": content_fingerprint(content),
                }
            )
    return normalized


def _chunk_documents_from_record(
    title: str,
    content: str,
    source: str,
    original_filename: str = "",
    document_id: str = "",
    segments: list[dict[str, Any]] | None = None,
) -> list[Document]:
    """Create chunked documents from one record."""

    documents: list[Document] = []
    normalized_segments = _normalize_segments(segments or [])
    chunk_units: list[tuple[str, dict[str, Any]]] = []
    if normalized_segments:
        for segment in normalized_segments:
            for chunk in split_text(
                segment["text"],
                chunk_size=settings.knowledge_chunk_size,
                overlap=settings.knowledge_chunk_overlap,
            ):
                location = {
                    key: segment[key]
                    for key in ("page", "section")
                    if segment.get(key) not in {None, ""}
                }
                chunk_units.append((chunk, location))
    else:
        chunk_units = [
            (chunk, {})
            for chunk in split_text(
                content,
                chunk_size=settings.knowledge_chunk_size,
                overlap=settings.knowledge_chunk_overlap,
            )
        ]
    stable_document_id = document_id or content_fingerprint(content)
    for index, (chunk, location) in enumerate(chunk_units):
        documents.append(
            Document(
                page_content=chunk,
                metadata={
                    "document_id": stable_document_id,
                    "title": title,
                    "source": source,
                    "original_filename": original_filename,
                    "chunk_index": index,
                    "chunk_id": f"{stable_document_id}:{index}",
                    **location,
                },
            )
        )
    return documents


def _build_documents() -> list[Document]:
    """Chunk all knowledge base records into documents."""

    documents: list[Document] = []
    for record in _load_knowledge_base_records():
        documents.extend(
            _chunk_documents_from_record(
                title=record["title"],
                content=record["content"],
                source=record["source"],
                original_filename=record["original_filename"],
                document_id=record["content_fingerprint"],
                segments=record.get("segments", []),
            )
        )
    return documents


def _chunk_ids(documents: list[Document]) -> list[str]:
    """Return stable IDs for vector-store persistence."""

    return [str(document.metadata["chunk_id"]) for document in documents]


def _embedding_profile() -> EmbeddingProfile:
    """Build the configured embedding profile without creating network clients."""

    return EmbeddingProfile(
        provider=settings.embedding_provider,  # type: ignore[arg-type]
        model=settings.embedding_model,
        dimensions=settings.embedding_dimensions,
        index_version=(
            f"{settings.embedding_index_version}-"
            f"chunk{settings.knowledge_chunk_size}-"
            f"overlap{settings.knowledge_chunk_overlap}"
        ),
        api_key=settings.dashscope_api_key,
        base_url=settings.dashscope_base_url,
    )


def _build_local_vectorstore(embeddings: Embeddings) -> _LocalVectorStore:
    """Create the fallback in-memory vector store."""

    return _LocalVectorStore(
        embeddings=embeddings,
        documents=_build_documents(),
    )


def _create_vectorstore_sync() -> VectorStoreLike:
    """Build or open the local persistent vector store."""

    profile = _embedding_profile()
    embeddings = create_embeddings(profile)
    if Chroma is None:
        return _build_local_vectorstore(embeddings)

    persist_dir = Path(settings.chroma_persist_dir)
    persist_dir.mkdir(parents=True, exist_ok=True)
    vectorstore = Chroma(
        collection_name=profile.collection_name(settings.chroma_collection_prefix),
        embedding_function=embeddings,
        persist_directory=str(persist_dir),
    )

    documents = _build_documents()
    stored = vectorstore.get(include=["documents", "metadatas"])
    stored_ids = [str(item) for item in stored.get("ids", [])]
    stored_documents = stored.get("documents", []) or []
    stored_metadatas = stored.get("metadatas", []) or []
    existing_by_id = {
        item_id: (page_content, metadata)
        for item_id, page_content, metadata in zip(
            stored_ids,
            stored_documents,
            stored_metadatas,
            strict=True,
        )
    }
    existing_ids = set(existing_by_id)
    expected_ids = set(_chunk_ids(documents))
    stale_ids = sorted(existing_ids - expected_ids)
    if stale_ids:
        vectorstore.delete(ids=stale_ids)
    changed_or_missing_documents = [
        document
        for document in documents
        if (
            str(document.metadata["chunk_id"]) not in existing_by_id
            or existing_by_id[str(document.metadata["chunk_id"])][0]
            != document.page_content
            or existing_by_id[str(document.metadata["chunk_id"])][1]
            != document.metadata
        )
    ]
    if changed_or_missing_documents:
        vectorstore.add_documents(
            changed_or_missing_documents,
            ids=_chunk_ids(changed_or_missing_documents),
        )
    return vectorstore


async def ensure_vectorstore() -> VectorStoreLike:
    """Create the vector store without blocking the event loop."""

    ensure_data_directories()
    vectorstore = await asyncio.to_thread(_create_vectorstore_sync)
    set_vectorstore(vectorstore)
    set_retriever(_build_retrieval_engine(vectorstore))
    return vectorstore


def _build_retrieval_engine(vectorstore: VectorStoreLike) -> RetrievalEngine:
    """Build one worker's lexical/fusion state from the durable knowledge JSON."""

    reranker = None
    if settings.reranker_enabled and settings.dashscope_api_key:
        reranker = DashScopeReranker(
            api_key=settings.dashscope_api_key,
            model=settings.reranker_model,
            endpoint=settings.reranker_endpoint,
            api_style=settings.reranker_api_style,  # type: ignore[arg-type]
            timeout_seconds=settings.reranker_timeout_seconds,
            max_document_chars=settings.reranker_max_document_chars,
        )
    return RetrievalEngine(
        vectorstore,
        _build_documents(),
        config=RetrievalConfig(
            dense_candidate_k=settings.retrieval_dense_candidate_k,
            lexical_candidate_k=settings.retrieval_lexical_candidate_k,
            rerank_candidate_k=settings.retrieval_rerank_candidate_k,
            rrf_k=settings.retrieval_rrf_k,
            dense_weight=settings.retrieval_dense_weight,
            lexical_weight=settings.retrieval_lexical_weight,
            production_strategy=settings.retrieval_strategy,  # type: ignore[arg-type]
        ),
        reranker=reranker,
    )


def search_knowledge_base_text(query: str, top_k: int = 4) -> str:
    """Search the enterprise knowledge base and return a compact summary."""

    try:
        retriever = get_retriever()
    except RuntimeError:
        return "知识库尚未初始化。"

    try:
        docs = retriever.retrieve(query, top_k=top_k).documents
    except Exception as exc:
        return f"知识库检索失败：{exc}"

    if not docs:
        return "未找到匹配的知识库内容。"

    lines: list[str] = []
    for idx, doc in enumerate(docs, start=1):
        title = str(doc.metadata.get("title", "未命名条目"))
        source = str(doc.metadata.get("source", "unknown"))
        snippet = doc.page_content.replace("\n", " ").strip()
        lines.append(f"{idx}. {title} [{source}]：{snippet[:200]}")
    return "\n".join(lines)


async def search_documents(query: str, top_k: int = 4) -> list[Document]:
    """Retrieve the most relevant documents from the vector store."""

    retriever = get_retriever()
    result = await asyncio.to_thread(retriever.retrieve, query, top_k)
    return result.documents


def keyword_overlap_score(query: str, texts: Iterable[str]) -> float:
    """Return a simple lexical overlap score."""

    query_tokens = set(tokenize(query))
    if not query_tokens:
        return 0.0

    combined_tokens: set[str] = set()
    for text in texts:
        combined_tokens.update(tokenize(text))
    if not combined_tokens:
        return 0.0
    return len(query_tokens & combined_tokens) / len(query_tokens)


def latest_user_text(messages: list[Any]) -> str:
    """Extract the latest human message content."""

    for message in reversed(messages):
        if getattr(message, "type", "") == "human":
            content = getattr(message, "content", "")
            return content if isinstance(content, str) else str(content)
    return ""


def extract_text_from_upload(filename: str, content: bytes) -> str:
    """Backward-compatible text-only view of structured extraction."""

    return extract_document_from_upload(filename, content).content


def extract_document_from_upload(filename: str, content: bytes) -> ExtractedDocument:
    """Extract text while retaining PDF pages and document section headings."""

    suffix = Path(filename).suffix.lower()
    segments: list[dict[str, Any]]
    if suffix in {".txt", ".md", ".py", ".log"}:
        text = content.decode("utf-8", errors="ignore")
        segments = _plain_text_segments(text, preserve_headings=suffix == ".md")
    elif suffix == ".csv":
        raw_text = content.decode("utf-8", errors="ignore")
        rows = csv.reader(StringIO(raw_text))
        text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        segments = [{"text": text}]
    elif suffix == ".json":
        obj = json.loads(content.decode("utf-8", errors="ignore"))
        text = json.dumps(obj, ensure_ascii=False, indent=2)
        segments = [{"text": text}]
    elif suffix == ".pdf":
        if PdfReader is None:
            raise ValueError("当前环境未安装 pypdf，无法解析 .pdf 文件。")
        reader = PdfReader(BytesIO(content))
        if len(reader.pages) > settings.max_document_pages:
            raise ValueError(f"PDF 页数不能超过 {settings.max_document_pages} 页。")
        segments = []
        extracted_length = 0
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            extracted_length += len(page_text)
            _check_extracted_length(extracted_length)
            if page_text:
                segments.append({"text": page_text, "page": page_number})
    elif suffix == ".docx":
        if DocxDocument is None:
            raise ValueError("当前环境未安装 python-docx，无法解析 .docx 文件。")
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                uncompressed_size = sum(item.file_size for item in archive.infolist())
        except zipfile.BadZipFile as exc:
            raise ValueError("DOCX 文件结构无效。") from exc
        if uncompressed_size > settings.max_archive_uncompressed_bytes:
            raise ValueError("DOCX 解压后的内容过大。")
        segments = _docx_segments(DocxDocument(BytesIO(content)))
    else:
        raise ValueError("仅支持 .txt .md .csv .json .pdf .docx .py .log 文件。")

    normalized = _normalize_segments(segments)
    text = "\n\n".join(segment["text"] for segment in normalized)
    _check_extracted_length(len(text))
    return ExtractedDocument(content=text, segments=normalized)


def _plain_text_segments(text: str, *, preserve_headings: bool) -> list[dict[str, Any]]:
    if not preserve_headings:
        return [{"text": text}]
    segments: list[dict[str, Any]] = []
    section = ""
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            item: dict[str, Any] = {"text": body}
            if section:
                item["section"] = section
            segments.append(item)
        buffer.clear()

    for line in text.splitlines():
        heading = line.strip()
        if heading.startswith("#") and heading.lstrip("#").strip():
            flush()
            section = heading.lstrip("#").strip()
            buffer.append(heading)
        else:
            buffer.append(line)
    flush()
    return segments or [{"text": text}]


def _docx_segments(document: Any) -> list[dict[str, Any]]:
    segments: list[dict[str, Any]] = []
    section = ""
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            indexed_text = f"{section}\n{body}" if section else body
            item: dict[str, Any] = {"text": indexed_text}
            if section:
                item["section"] = section
            segments.append(item)
        buffer.clear()

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()
        style_name = str(getattr(paragraph.style, "name", ""))
        if paragraph_text and style_name.casefold().startswith("heading"):
            flush()
            section = paragraph_text
        elif paragraph_text:
            buffer.append(paragraph_text)
    flush()
    for table_number, table in enumerate(document.tables, start=1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text:
            segments.append({"text": table_text, "section": f"Table {table_number}"})
    _check_extracted_length(sum(len(item["text"]) for item in segments))
    return segments


def _check_extracted_length(length: int) -> None:
    if length > settings.max_extracted_chars:
        raise ValueError("文档解析后的文本过大。")


def save_uploaded_file(filename: str, content: bytes) -> Path:
    """Persist the uploaded original file locally."""

    ensure_data_directories()
    safe_filename = Path(filename.replace("\\", "/")).name.replace("\x00", "")
    if safe_filename in {"", ".", ".."}:
        raise ValueError("上传文件名无效。")
    source = Path(safe_filename)
    handle, target_name = tempfile.mkstemp(
        prefix=f"{source.stem[:80]}-",
        suffix=source.suffix,
        dir=settings.upload_dir,
    )
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
    except Exception:
        Path(target_name).unlink(missing_ok=True)
        raise
    return Path(target_name)


def add_knowledge_record(
    title: str,
    content: str,
    source: str,
    original_filename: str,
    segments: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Serialize one knowledge write across JSON, vector, and lexical indexes."""

    lock_path = f"{settings.knowledge_base_path}.lock"
    with _KNOWLEDGE_WRITE_LOCK, FileLock(
        lock_path,
        timeout=settings.knowledge_write_lock_timeout_seconds,
    ):
        return _add_knowledge_record_unlocked(
            title,
            content,
            source,
            original_filename,
            segments,
        )


def _add_knowledge_record_unlocked(
    title: str,
    content: str,
    source: str,
    original_filename: str,
    segments: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Append one record to the knowledge base and index it immediately."""

    clean_title = title.strip() or Path(original_filename).stem
    clean_content = content.strip()
    if not clean_content:
        raise ValueError("上传文件解析后内容为空，无法入库。")

    kb_path = Path(settings.knowledge_base_path)
    records = _load_json_file(kb_path)
    candidate_fingerprint = content_fingerprint(clean_content)
    normalized_segments = _normalize_segments(segments or [])
    duplicate = find_duplicate(
        clean_content,
        records=records,
        near_threshold=settings.knowledge_near_duplicate_threshold,
        minimum_length_ratio=settings.knowledge_near_duplicate_min_length_ratio,
        minimum_near_length=settings.knowledge_near_duplicate_min_length,
    )
    if duplicate is not None:
        existing = duplicate.record
        existing_content = str(existing.get("content", "")).strip()
        metadata_upgraded = False
        if (
            duplicate.kind == "exact"
            and normalized_segments
            and not _normalize_segments(existing.get("segments", []))
        ):
            existing = _upgrade_duplicate_segments(
                records=records,
                existing=existing,
                segments=normalized_segments,
                kb_path=kb_path,
            )
            metadata_upgraded = True
        return {
            "title": str(existing.get("title", "")).strip() or clean_title,
            "content": existing_content,
            "source": str(existing.get("source", "seed")).strip() or "seed",
            "original_filename": str(existing.get("original_filename", "")).strip(),
            "deduplicated": True,
            "dedup_type": "exact" if duplicate.kind == "exact" else "similar",
            "similarity": duplicate.score,
            "metadata_upgraded": metadata_upgraded,
        }

    record = {
        "title": clean_title,
        "content": clean_content,
        "source": source,
        "original_filename": original_filename,
        "content_fingerprint": candidate_fingerprint,
        "deduplicated": False,
        "dedup_type": "none",
    }
    if normalized_segments:
        record["segments"] = normalized_segments
    documents = _chunk_documents_from_record(
        title=clean_title,
        content=clean_content,
        source=source,
        original_filename=original_filename,
        document_id=candidate_fingerprint,
        segments=normalized_segments,
    )
    vectorstore = get_vectorstore()
    document_ids = _chunk_ids(documents)
    vectorstore.add_documents(
        documents,
        ids=document_ids,
    )
    records.append(record)
    try:
        _save_json_file(kb_path, records)
    except Exception as persistence_error:
        try:
            vectorstore.delete(ids=document_ids)
        except Exception as rollback_error:
            persistence_error.add_note(f"Vector rollback also failed: {rollback_error}")
        raise
    get_retriever().add_documents(documents)
    return record


def _upgrade_duplicate_segments(
    *,
    records: list[dict[str, Any]],
    existing: dict[str, Any],
    segments: list[dict[str, Any]],
    kb_path: Path,
) -> dict[str, Any]:
    """Backfill locations for an exact legacy record and rebuild its chunks."""

    index = next(i for i, record in enumerate(records) if record is existing)
    updated = dict(existing)
    updated["segments"] = segments
    title = str(existing.get("title", "")).strip()
    content = str(existing.get("content", "")).strip()
    source = str(existing.get("source", "seed")).strip() or "seed"
    original_filename = str(existing.get("original_filename", "")).strip()
    document_id = content_fingerprint(content)
    old_documents = _chunk_documents_from_record(
        title,
        content,
        source,
        original_filename,
        document_id,
    )
    new_documents = _chunk_documents_from_record(
        title,
        content,
        source,
        original_filename,
        document_id,
        segments,
    )
    old_ids = set(_chunk_ids(old_documents))
    new_ids = set(_chunk_ids(new_documents))
    vectorstore = get_vectorstore()
    vectorstore.add_documents(new_documents, ids=_chunk_ids(new_documents))
    stale_ids = sorted(old_ids - new_ids)
    if stale_ids:
        vectorstore.delete(ids=stale_ids)
    records[index] = updated
    try:
        _save_json_file(kb_path, records)
    except Exception as persistence_error:
        try:
            vectorstore.delete(ids=sorted(new_ids))
            vectorstore.add_documents(old_documents, ids=_chunk_ids(old_documents))
        except Exception as rollback_error:
            persistence_error.add_note(f"Vector rollback also failed: {rollback_error}")
        raise
    set_retriever(_build_retrieval_engine(vectorstore))
    return updated


def list_knowledge_records() -> list[dict[str, Any]]:
    """Return a compact list of knowledge records."""

    records = _load_json_file(Path(settings.knowledge_base_path))
    items: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=1):
        content = str(record.get("content", "")).strip()
        items.append(
            {
                "id": index,
                "title": str(record.get("title", "")).strip(),
                "source": str(record.get("source", "seed")).strip() or "seed",
                "original_filename": str(record.get("original_filename", "")).strip(),
                "preview": content[:120],
                "content_fingerprint": str(record.get("content_fingerprint", "")).strip(),
            }
        )
    return items
